"""Construit « MemoBook · Tableau de bord.xlsx », à importer dans Google Drive.

Les onglets « Relevés », « Résiliations » et « Carnets à livrer » portent
exactement les en-têtes que `StatsSheet.gs` écrit : le script les retrouve
remplis d'en-têtes, n'en recrée pas, et pose ses lignes en dessous.
L'onglet « Tableau de bord » ne contient que des formules sur ces trois-là.

    python3 -m venv .venv && .venv/bin/pip install openpyxl
    .venv/bin/python build-dashboard-workbook.py "MemoBook - Tableau de bord.xlsx"

Les formules restent en syntaxe anglaise (virgules) : Google Sheets les
traduit à l'import. Les recherches sont bornées à 1000 lignes, comme les
courbes — un relevé par jour, c'est presque trois ans.
"""

import sys
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = sys.argv[1]

# --- Couleurs MemoBook (ios/Modules/Sources/MemoBookDesign/Tokens.swift) ---
GREEN = "28654B"
GREEN_LIGHT = "3D9A6F"
INK = "2D231A"
BACKGROUND = "FCF2E9"
SURFACE = "FFFCF8"
BLUE = "AFD2F0"
BLUE_TEXT = "4088C6"
ACCENT = "E2F32B"
SEPARATOR = "CFBBAA"
BEIGE = "F9E6D6"
BEIGE_NOTICE = "E6D5C4"
VALID = "3FA673"
WARNING = "FF682C"
ERROR = "DE2B2E"
GREY = "8E8E93"

FAMILY = "Poppins"
MAX_ROWS = 1000  # lignes pré-formatées sous chaque en-tête


def font(size=10, bold=False, color=INK):
    return Font(name=FAMILY, size=size, bold=bold, color=color)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


HAIRLINE = Side(style="thin", color=SEPARATOR)
BOX = Border(left=HAIRLINE, right=HAIRLINE, top=HAIRLINE, bottom=HAIRLINE)
UNDERLINE = Border(bottom=Side(style="medium", color=GREEN))

FORMAT_DATE = "dd/mm/yyyy hh:mm"
FORMAT_DAY = "dd/mm/yyyy"
FORMAT_INT = "#,##0"
FORMAT_EUR = '#,##0.00 "€"'
FORMAT_DELTA = '+#,##0;-#,##0;"="'
FORMAT_DELTA_EUR = '+#,##0 "€";-#,##0 "€";"="'

# --- Ce que le script écrit --------------------------------------------------

SNAPSHOT_COLUMNS = [
    ("Relevé", FORMAT_DATE, 18),
    ("Utilisateurs", FORMAT_INT, 13),
    ("Nouveaux (7 j)", FORMAT_INT, 14),
    ("Abonnés", FORMAT_INT, 11),
    ("Voyages", FORMAT_INT, 11),
    ("Carnets en cours", FORMAT_INT, 16),
    ("Voyages à venir", FORMAT_INT, 15),
    ("Voyages passés", FORMAT_INT, 15),
    ("Carnets composés", FORMAT_INT, 17),
    ("Carnets commandés", FORMAT_INT, 18),
    ("Commandes en cours", FORMAT_INT, 19),
    ("Brouillons", FORMAT_INT, 12),
    ("Soumises", FORMAT_INT, 11),
    ("En production", FORMAT_INT, 14),
    ("Expédiées", FORMAT_INT, 11),
    ("Livrées", FORMAT_INT, 10),
    ("Annulées", FORMAT_INT, 11),
    ("Exemplaires", FORMAT_INT, 13),
    ("Chiffre d'affaires (€)", FORMAT_EUR, 20),
    ("Abonnements actifs", FORMAT_INT, 19),
    ("Abonnements résiliés", FORMAT_INT, 20),
    ("Abonnements arrivés à terme", FORMAT_INT, 26),
    ("Souvenirs racontés", FORMAT_INT, 18),
    ("Photos", FORMAT_INT, 10),
]

CANCELLATION_COLUMNS = [
    ("Date", FORMAT_DATE, 18),
    ("Compte", "@", 12),
    ("Raison", "@", 44),
    ("Abonnement actif", "@", 17),
]
REASON_SUMMARY_COLUMNS = [("Raison", "@", 44), ("Total", FORMAT_INT, 10)]  # F:G

DELIVERY_COLUMNS = [
    ("Commande", "@", 30),
    ("Voyage", "@", 28),
    ("Statut", "@", 14),
    ("Exemplaires", FORMAT_INT, 12),
    ("Pages", FORMAT_INT, 9),
    ("Montant (€)", FORMAT_EUR, 13),
    ("Destinataire", "@", 24),
    ("Ville", "@", 18),
    ("Pays", "@", 10),
    ("Suivi", "@", 40),
    ("Soumise le", FORMAT_DATE, 18),
    ("Expédiée le", FORMAT_DATE, 18),
]


def data_sheet(wb, title, columns, first_col=1, banded=True):
    """Un onglet de données : en-tête vert, lignes pré-formatées en Poppins."""
    ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
    ws.sheet_properties.tabColor = GREEN
    ws.sheet_view.showGridLines = False
    for offset, (header, number_format, width) in enumerate(columns):
        col = first_col + offset
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        head = ws.cell(row=1, column=col, value=header)
        head.font = font(10, bold=True, color=SURFACE)
        head.fill = fill(GREEN)
        head.alignment = Alignment(vertical="center", wrap_text=True)
        head.border = BOX
        for row in range(2, MAX_ROWS + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = font(10)
            cell.number_format = number_format
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="hair", color=SEPARATOR))
            if banded and row % 2 == 0:
                cell.fill = fill(SURFACE)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = ws.cell(row=2, column=first_col + 1 if first_col == 1 else 1)
    return ws


# --- Le tableau de bord ------------------------------------------------------

def col_of(header):
    """La lettre de colonne d'un chiffre dans « Relevés »."""
    for index, (name, _, _) in enumerate(SNAPSHOT_COLUMNS, start=1):
        if name == header:
            return get_column_letter(index)
    raise KeyError(header)


LAST_DATE = f"INDEX(Relevés!$A$1:$A${MAX_ROWS},COUNTA(Relevés!$A:$A))"


def latest(header):
    """La dernière valeur relevée d'une colonne. Vide tant qu'aucun relevé."""
    c = col_of(header)
    return f'=IFERROR(IF(COUNTA(Relevés!$A:$A)<2,"",INDEX(Relevés!${c}$1:${c}${MAX_ROWS},COUNTA(Relevés!${c}:${c}))),"")'


def week_ago(header):
    """La valeur relevée sept jours plus tôt (ou le relevé juste avant).

    Les relevés sont rangés par date : compter ceux d'avant la date visée
    donne la ligne à lire. Pas de MATCH approché sur des blancs."""
    c = col_of(header)
    count = f'COUNTIF(Relevés!$A$2:$A${MAX_ROWS},"<="&({LAST_DATE}-7))'
    return f'=IFERROR(IF({count}=0,"",INDEX(Relevés!${c}$2:${c}${MAX_ROWS},{count})),"")'


SECTIONS = [
    (
        "Utilisateurs",
        [
            ("Utilisateurs", "Comptes ouverts", FORMAT_INT, FORMAT_DELTA),
            ("Nouveaux (7 j)", "Nouveaux sur 7 jours", FORMAT_INT, FORMAT_DELTA),
            ("Abonnés", "Abonnés", FORMAT_INT, FORMAT_DELTA),
            ("Abonnements résiliés", "Résiliations", FORMAT_INT, FORMAT_DELTA),
        ],
    ),
    (
        "Voyages",
        [
            ("Voyages", "Voyages créés", FORMAT_INT, FORMAT_DELTA),
            ("Carnets en cours", "Carnets en cours", FORMAT_INT, FORMAT_DELTA),
            ("Voyages à venir", "Voyages à venir", FORMAT_INT, FORMAT_DELTA),
            ("Carnets composés", "Carnets composés", FORMAT_INT, FORMAT_DELTA),
        ],
    ),
    (
        "Carnets",
        [
            ("Carnets commandés", "Carnets commandés", FORMAT_INT, FORMAT_DELTA),
            ("Commandes en cours", "Commandes en cours", FORMAT_INT, FORMAT_DELTA),
            ("Livrées", "Carnets livrés", FORMAT_INT, FORMAT_DELTA),
            ("Chiffre d'affaires (€)", "Chiffre d'affaires", FORMAT_EUR, FORMAT_DELTA_EUR),
        ],
    ),
    (
        "Souvenirs",
        [
            ("Souvenirs racontés", "Souvenirs racontés", FORMAT_INT, FORMAT_DELTA),
            ("Photos", "Photos", FORMAT_INT, FORMAT_DELTA),
            ("Exemplaires", "Exemplaires imprimés", FORMAT_INT, FORMAT_DELTA),
            ("Abonnements actifs", "Abonnements actifs", FORMAT_INT, FORMAT_DELTA),
        ],
    ),
]

TILE_WIDTH = 3  # colonnes par tuile
FIRST_TILE_COL = 2  # B


def dashboard(wb):
    ws = wb.active
    ws.title = "Tableau de bord"
    ws.sheet_properties.tabColor = ACCENT
    ws.sheet_view.showGridLines = False

    # Le fond crème de l'app, partout.
    for row in range(1, 80):
        for col in range(1, 20):
            ws.cell(row=row, column=col).fill = fill(BACKGROUND)
    ws.column_dimensions["A"].width = 3
    for col in range(2, 20):
        ws.column_dimensions[get_column_letter(col)].width = 7.5

    # Bandeau
    ws.merge_cells("B2:N2")
    ws["B2"] = "MemoBook · Tableau de bord"
    ws["B2"].font = font(22, bold=True, color=SURFACE)
    ws["B2"].fill = fill(GREEN)
    ws["B2"].alignment = Alignment(vertical="center", indent=1)
    for col in range(2, 15):
        ws.cell(row=2, column=col).fill = fill(GREEN)
    ws.row_dimensions[2].height = 44

    ws["B3"] = "Dernier relevé"
    ws["B3"].font = font(9, color=GREY)
    ws.merge_cells("B4:E4")
    ws["B4"] = f'=IFERROR(IF(COUNTA(Relevés!$A:$A)<2,"Aucun relevé — lance npm run stats:push",{LAST_DATE}),"")'
    ws["B4"].number_format = 'dddd d mmmm yyyy "à" hh:mm'
    ws["B4"].font = font(11, bold=True, color=GREEN)

    ws["G3"] = "Carnets à livrer"
    ws["G3"].font = font(9, color=GREY)
    ws["G4"] = "=MAX(0,COUNTA('Carnets à livrer'!$A:$A)-1)"
    ws["G4"].number_format = FORMAT_INT
    ws["G4"].font = font(11, bold=True, color=WARNING)

    ws["J3"] = "Raisons de départ reçues"
    ws["J3"].font = font(9, color=GREY)
    ws["J4"] = "=MAX(0,COUNTA(Résiliations!$A:$A)-1)"
    ws["J4"].number_format = FORMAT_INT
    ws["J4"].font = font(11, bold=True, color=BLUE_TEXT)

    row = 6
    for section_title, tiles in SECTIONS:
        ws.cell(row=row, column=FIRST_TILE_COL, value=section_title).font = font(12, bold=True, color=GREEN)
        for col in range(FIRST_TILE_COL, FIRST_TILE_COL + len(tiles) * (TILE_WIDTH + 1) - 1):
            ws.cell(row=row, column=col).border = UNDERLINE
        row += 1
        for index, (header, label, number_format, delta_format) in enumerate(tiles):
            c0 = FIRST_TILE_COL + index * (TILE_WIDTH + 1)
            c1 = c0 + TILE_WIDTH - 1
            a, b = get_column_letter(c0), get_column_letter(c1)
            for r in range(row, row + 3):
                for c in range(c0, c1 + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.fill = fill(SURFACE)
                    cell.border = Border(
                        left=HAIRLINE if c == c0 else None,
                        right=HAIRLINE if c == c1 else None,
                        top=HAIRLINE if r == row else None,
                        bottom=HAIRLINE if r == row + 2 else None,
                    )
            ws.merge_cells(f"{a}{row}:{b}{row}")
            ws[f"{a}{row}"] = label
            ws[f"{a}{row}"].font = font(9, color=GREY)
            ws[f"{a}{row}"].alignment = Alignment(indent=1, vertical="center")

            ws.merge_cells(f"{a}{row + 1}:{b}{row + 1}")
            ws[f"{a}{row + 1}"] = latest(header)
            ws[f"{a}{row + 1}"].number_format = number_format
            ws[f"{a}{row + 1}"].font = font(20, bold=True, color=INK)
            ws[f"{a}{row + 1}"].alignment = Alignment(indent=1, vertical="center")

            ws.merge_cells(f"{a}{row + 2}:{b}{row + 2}")
            ws[f"{a}{row + 2}"] = (
                f'=IFERROR(IF(OR({a}{row + 1}="",{week_ago(header)[1:]}=""),"",'
                f'{a}{row + 1}-{week_ago(header)[1:]}),"")'
            )
            ws[f"{a}{row + 2}"].number_format = delta_format
            ws[f"{a}{row + 2}"].font = font(9, color=GREEN_LIGHT)
            ws[f"{a}{row + 2}"].alignment = Alignment(indent=1, vertical="center")
        ws.row_dimensions[row].height = 18
        ws.row_dimensions[row + 1].height = 34
        ws.row_dimensions[row + 2].height = 18
        row += 4

    ws.cell(row=row - 1, column=FIRST_TILE_COL, value="La petite ligne sous chaque chiffre : l'écart avec le relevé d'il y a sept jours.").font = font(8, color=GREY)

    # Raisons de départ (le cumul que le script pose en Résiliations!F:G)
    row += 1
    ws.cell(row=row, column=FIRST_TILE_COL, value="Pourquoi ils partent").font = font(12, bold=True, color=GREEN)
    for col in range(FIRST_TILE_COL, FIRST_TILE_COL + 7):
        ws.cell(row=row, column=col).border = UNDERLINE
    reasons_first = row + 1
    for i in range(1, 9):
        r = row + i
        ws.merge_cells(start_row=r, start_column=FIRST_TILE_COL, end_row=r, end_column=FIRST_TILE_COL + 5)
        cell = ws.cell(row=r, column=FIRST_TILE_COL, value=f'=IFERROR(IF(Résiliations!$F${i + 1}="","",Résiliations!$F${i + 1}),"")')
        cell.font = font(10)
        cell.alignment = Alignment(indent=1)
        total = ws.cell(row=r, column=FIRST_TILE_COL + 6, value=f'=IFERROR(IF(Résiliations!$G${i + 1}="","",Résiliations!$G${i + 1}),"")')
        total.font = font(10, bold=True, color=INK)
        total.number_format = FORMAT_INT
        total.alignment = Alignment(horizontal="right")
        for c in range(FIRST_TILE_COL, FIRST_TILE_COL + 7):
            ws.cell(row=r, column=c).fill = fill(SURFACE if i % 2 else BEIGE)

    # Graphiques sur Relevés
    charts_row = 6
    charts_col = "R"
    ws.column_dimensions["Q"].width = 3
    for title, headers, anchor_row in [
        ("Utilisateurs et abonnés", ["Utilisateurs", "Abonnés", "Abonnements actifs"], charts_row),
        ("Carnets", ["Carnets commandés", "Commandes en cours", "Livrées", "Carnets composés"], charts_row + 16),
        ("Chiffre d'affaires (€)", ["Chiffre d'affaires (€)"], charts_row + 32),
    ]:
        chart = LineChart()
        chart.title = title
        chart.style = 2
        chart.height = 7.5
        chart.width = 16
        chart.y_axis.majorGridlines = None
        chart.legend.position = "b"
        data_ws = wb["Relevés"]
        for header in headers:
            c = SNAPSHOT_COLUMNS.index(next(x for x in SNAPSHOT_COLUMNS if x[0] == header)) + 1
            chart.add_data(Reference(data_ws, min_col=c, min_row=1, max_row=MAX_ROWS), titles_from_data=True)
        chart.set_categories(Reference(data_ws, min_col=1, min_row=2, max_row=MAX_ROWS))
        chart.x_axis.number_format = FORMAT_DAY
        palette = [GREEN, BLUE_TEXT, WARNING, GREEN_LIGHT]
        for series, color in zip(chart.series, palette):
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 22000
            series.smooth = False
        ws.add_chart(chart, f"{charts_col}{anchor_row}")

    return ws


def guide(wb):
    ws = wb.create_sheet("Mode d'emploi")
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    for row in range(1, 40):
        ws.cell(row=row, column=1).fill = fill(BACKGROUND)
        ws.cell(row=row, column=2).fill = fill(BACKGROUND)
        ws.cell(row=row, column=3).fill = fill(BACKGROUND)

    lines = [
        ("Brancher la feuille de bord", 16, True, GREEN),
        ("", 6, False, INK),
        ("L'API MemoBook n'écrit pas ici directement : elle envoie un JSON à un script Apps Script, et c'est lui qui pose les lignes. Sept étapes, une seule fois.", 10, False, INK),
        ("", 6, False, INK),
        ("1.  Importe ce classeur dans Google Drive : Nouveau → Importer un fichier, puis ouvre-le avec Google Sheets et enregistre-le au format Sheets (Fichier → Enregistrer au format Google Sheets).", 10, False, INK),
        ("2.  Extensions → Apps Script. Remplace le contenu de Code.gs par le fichier StatsSheet.gs, enregistre.", 10, False, INK),
        ("3.  Dans Apps Script : Paramètres du projet → Propriétés du script → ajoute SECRET avec une longue chaîne au hasard.", 10, False, INK),
        ("4.  Déployer → Nouveau déploiement → type « Application web » : exécuter en tant que « Moi », accès « Tout le monde ». Le secret filtre les appels.", 10, False, INK),
        ("5.  Copie l'URL du déploiement (…/exec) dans STATS_SHEET_WEBHOOK_URL sur Railway, et le secret dans STATS_SHEET_SECRET. Redéploie l'API.", 10, False, INK),
        ("6.  Depuis backend/ : npm run stats:push. Le premier relevé arrive dans « Relevés », les carnets à livrer dans leur onglet, et le tableau de bord se remplit.", 10, False, INK),
        ("7.  Ensuite, l'API envoie un relevé chaque nuit à 4 h 20 UTC, et chaque raison de départ à l'instant où elle est donnée.", 10, False, INK),
        ("", 6, False, INK),
        ("Les onglets", 13, True, GREEN),
        ("Tableau de bord — des formules sur les trois autres onglets, rien à saisir. Les graphiques lisent « Relevés » jusqu'à la ligne 1000.", 10, False, INK),
        ("Relevés — une ligne par jour. Deux relevés le même jour : le second remplace le premier. Ne pas renommer les colonnes : le script les cherche par position.", 10, False, INK),
        ("Résiliations — colonnes A à D, une ligne par départ ; colonnes F et G, le cumul par raison, réécrit à chaque relevé.", 10, False, INK),
        ("Carnets à livrer — réécrit en entier à chaque relevé : les commandes soumises, en production ou expédiées et pas encore livrées.", 10, False, INK),
        ("", 6, False, INK),
        ("Vérifier que le script répond", 13, True, GREEN),
        ("Ouvre l'URL du déploiement dans un navigateur : elle doit afficher {\"ok\":true,\"message\":\"MemoBook · feuille de bord\"}. Une nouvelle version du script demande un nouveau déploiement (Déployer → Gérer les déploiements → modifier → nouvelle version).", 10, False, INK),
    ]
    for index, (text, size, bold, color) in enumerate(lines, start=2):
        cell = ws.cell(row=index, column=2, value=text)
        cell.font = font(size, bold=bold, color=color)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if size == 10 and text:
            ws.row_dimensions[index].height = 30 if len(text) > 120 else 18
    return ws


wb = Workbook()
wb.active.title = "Tableau de bord"
data_sheet(wb, "Relevés", SNAPSHOT_COLUMNS)
cancel = data_sheet(wb, "Résiliations", CANCELLATION_COLUMNS)
data_sheet(wb, "Résiliations", REASON_SUMMARY_COLUMNS, first_col=6, banded=False)
cancel.column_dimensions["E"].width = 3
cancel.freeze_panes = "A2"
data_sheet(wb, "Carnets à livrer", DELIVERY_COLUMNS)
dashboard(wb)
guide(wb)
wb.move_sheet("Mode d'emploi", offset=-3)  # juste après le tableau de bord
wb.active = 0
wb.save(OUT)
print("écrit", OUT)
